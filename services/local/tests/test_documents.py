"""文档解析与导出测试：TXT / Markdown / DOCX / PDF（坐标）/ SRT 与 MD/DOCX 导出。"""

from __future__ import annotations

import base64
import io

import pytest
from fastapi.testclient import TestClient

from .conftest import FIXTURES_DIR, b64, make_docx_bytes, make_pdf_bytes, make_wav_bytes


def _parse(client: TestClient, payload: bytes, **kwargs):
    request = {"requestId": "req-doc-1", "data": b64(payload)}
    request.update(kwargs)
    return client.post("/v1/documents/parse", json=request)


def _export(client: TestClient, units, **kwargs):
    request = {"requestId": "req-doc-export-1", "units": units}
    request.update(kwargs)
    return client.post("/v1/documents/export", json=request)


# ---------------------------------------------------------------- TXT
def test_parse_txt_fixture(authed_client: TestClient):
    sample = (FIXTURES_DIR / "documents" / "sample.txt").read_bytes()
    response = _parse(authed_client, sample, filename="sample.txt", mimeType="text/plain")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["format"] == "txt"
    assert body["unitCount"] == 3
    assert [unit["kind"] for unit in body["units"]] == ["paragraph"] * 3
    assert body["units"][0]["text"].startswith("Hello world.")
    assert body["requiresOcr"] is False


# ---------------------------------------------------------------- Markdown
def test_parse_markdown_structure(authed_client: TestClient):
    markdown = (
        "# Title\n\n"
        "First paragraph text.\n\n"
        "## Section\n\n"
        "- item one\n"
        "- item two\n\n"
        "```python\nprint('hi')\n```\n"
    ).encode("utf-8")
    response = _parse(authed_client, markdown, filename="readme.md", mimeType="text/markdown")
    assert response.status_code == 200, response.text
    body = response.json()
    kinds = [unit["kind"] for unit in body["units"]]
    assert kinds == ["heading", "paragraph", "heading", "list-item", "list-item", "code"]
    assert body["units"][0]["level"] == 1
    assert body["units"][2]["level"] == 2
    assert body["units"][-1]["translatable"] is False


# ---------------------------------------------------------------- DOCX
def test_parse_docx_units_and_order(authed_client: TestClient):
    payload = make_docx_bytes(
        paragraphs=[
            ("Document Title", "Title"),
            ("Introduction paragraph.", None),
            ("Chapter One", "Heading 1"),
            ("Body paragraph.", None),
            ("Bullet item", "List Bullet"),
        ],
        tables=[[["Name", "Value"], ["alpha", "1"]]],
    )
    response = _parse(
        authed_client,
        payload,
        filename="sample.docx",
        mimeType="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["format"] == "docx"
    assert body["title"] == "Document Title"

    kinds = [unit["kind"] for unit in body["units"]]
    assert kinds.count("heading") == 2  # Title + Heading 1
    assert kinds.count("paragraph") == 2
    assert kinds.count("list-item") == 1
    assert kinds.count("table-cell") == 4

    # 阅读顺序与文档顺序一致
    orders = [unit["readingOrder"] for unit in body["units"]]
    assert orders == sorted(orders)

    # 表格结构保留
    cells = [unit for unit in body["units"] if unit["kind"] == "table-cell"]
    assert cells[0]["tableId"]
    assert cells[0]["rowIndex"] == 0 and cells[0]["columnIndex"] == 0
    assert cells[0]["text"] == "Name"
    assert cells[-1]["text"] == "1"

    # 段落 ID 唯一
    ids = [unit["id"] for unit in body["units"]]
    assert len(ids) == len(set(ids))


def test_parse_docx_rejects_invalid_payload(authed_client: TestClient):
    response = _parse(authed_client, b"not a docx at all", filename="broken.docx")
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "INVALID_REQUEST"


# ---------------------------------------------------------------- PDF
def test_parse_pdf_with_coordinates(authed_client: TestClient):
    payload = make_pdf_bytes(
        [
            [("Hello PDF world.", (72, 100)), ("Second block.", (72, 200))],
            [("Page two content.", (72, 80))],
        ]
    )
    response = _parse(authed_client, payload, filename="sample.pdf", mimeType="application/pdf")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["format"] == "pdf"
    assert body["pages"] == 2
    assert body["requiresOcr"] is False

    units = body["units"]
    assert len(units) == 3
    # 阅读顺序：第 1 页两块，然后第 2 页
    assert units[0]["pageNumber"] == 1
    assert units[1]["pageNumber"] == 1
    assert units[2]["pageNumber"] == 2
    assert [unit["readingOrder"] for unit in units] == [0, 1, 2]
    # bbox = [x, y, width, height]
    assert len(units[0]["bbox"]) == 4
    assert units[0]["bbox"][0] >= 0 and units[0]["bbox"][1] >= 0
    assert units[0]["bbox"][2] > 0 and units[0]["bbox"][3] > 0
    assert "Hello PDF world." in units[0]["text"]


def test_parse_pdf_detects_scanned_document(authed_client: TestClient):
    """仅含图像、无文本层的 PDF 标记 requiresOcr。"""
    import fitz
    from PIL import Image

    image = Image.new("RGB", (300, 200), color=(255, 255, 255))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    png_bytes = buffer.getvalue()

    document = fitz.open()
    page = document.new_page(width=612, height=792)
    page.insert_image(fitz.Rect(50, 50, 350, 250), stream=png_bytes)
    pdf_buffer = io.BytesIO()
    document.save(pdf_buffer)
    payload = pdf_buffer.getvalue()
    document.close()

    response = _parse(authed_client, payload, filename="scanned.pdf", mimeType="application/pdf")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["requiresOcr"] is True
    assert any("document-ocr" in warning for warning in body["warnings"])


# ---------------------------------------------------------------- SRT
def test_parse_srt_preserves_timeline(authed_client: TestClient):
    sample = (FIXTURES_DIR / "subtitles" / "sample.srt").read_bytes()
    response = _parse(authed_client, sample, filename="sample.srt", mimeType="application/x-subrip")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["format"] == "srt"
    cues = body["units"]
    assert len(cues) == 2
    assert cues[0]["kind"] == "subtitle-cue"
    assert cues[0]["startTime"] == 1.0
    assert cues[0]["endTime"] == 3.5
    assert "Welcome" in cues[0]["text"]
    assert cues[1]["startTime"] == 4.0


# ---------------------------------------------------------------- 导出
def test_export_markdown_bilingual(authed_client: TestClient):
    units = [
        {"id": "u-1", "text": "Hello", "translation": "你好", "kind": "paragraph"},
        {"id": "u-2", "text": "World", "translation": "世界", "kind": "paragraph"},
    ]
    response = _export(authed_client, units, format="markdown", title="Demo")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["format"] == "markdown"
    assert body["filename"].endswith(".md")
    assert body["mimeType"] == "text/markdown"

    content = base64.b64decode(body["data"]).decode("utf-8")
    assert "# Demo" in content
    assert "你好" in content
    assert "> Hello" in content
    assert "世界" in content
    assert body["sizeBytes"] == len(content.encode("utf-8"))


def test_export_markdown_translation_only(authed_client: TestClient):
    units = [{"id": "u-1", "text": "Hello", "translation": "你好"}]
    response = _export(authed_client, units, format="markdown", layout="translation-only")
    assert response.status_code == 200
    content = base64.b64decode(response.json()["data"]).decode("utf-8")
    assert "你好" in content
    assert "> Hello" not in content


def test_export_docx_bilingual_roundtrip(authed_client: TestClient):
    units = [
        {"id": "u-1", "text": "Hello", "translation": "你好", "kind": "heading", "level": 1},
        {"id": "u-2", "text": "World", "translation": "世界", "kind": "paragraph"},
    ]
    response = _export(authed_client, units, format="docx", title="Demo")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["format"] == "docx"
    assert body["filename"].endswith(".docx")
    assert body["mimeType"].endswith("wordprocessingml.document")

    payload = base64.b64decode(body["data"])
    import docx

    document = docx.Document(io.BytesIO(payload))
    texts = [p.text for p in document.paragraphs if p.text.strip()]
    assert "Demo" in texts
    assert "你好" in texts
    assert "Hello" in texts  # 对照阅读：原文以注释段落保留
    assert "世界" in texts


def test_export_docx_table_reassembly(authed_client: TestClient):
    units = [
        {"id": "c-1", "text": "Name", "translation": "名称", "kind": "table-cell",
         "tableId": "tbl-1", "rowIndex": 0, "columnIndex": 0},
        {"id": "c-2", "text": "Value", "translation": "值", "kind": "table-cell",
         "tableId": "tbl-1", "rowIndex": 0, "columnIndex": 1},
        {"id": "c-3", "text": "alpha", "translation": "阿尔法", "kind": "table-cell",
         "tableId": "tbl-1", "rowIndex": 1, "columnIndex": 0},
        {"id": "c-4", "text": "1", "translation": "一", "kind": "table-cell",
         "tableId": "tbl-1", "rowIndex": 1, "columnIndex": 1},
    ]
    response = _export(authed_client, units, format="docx")
    assert response.status_code == 200, response.text
    payload = base64.b64decode(response.json()["data"])

    import docx

    document = docx.Document(io.BytesIO(payload))
    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.cell(0, 0).text == "名称"
    assert table.cell(0, 1).text == "值"
    assert table.cell(1, 0).text == "阿尔法"


def test_export_rejects_empty_units(authed_client: TestClient):
    response = _export(authed_client, [], format="markdown")
    assert response.status_code == 422


def test_parse_rejects_invalid_base64(authed_client: TestClient):
    response = authed_client.post(
        "/v1/documents/parse",
        json={"requestId": "req-bad", "data": "!!!not-base64!!!"},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "INVALID_REQUEST"
